"""Export engine (V2 canonical export + templates).

Builds the canonical model straight from the database and serializes it back
into the exact shapes the V2 import pipeline consumes, so that:

    export -> import  (CSV or XLSX)

round-trips symmetrically. The canonical model is the single source of truth:
bills carry their charges, derived fields (amountReceived / paymentStatus / Total)
are recomputed before serialization, and payments are the only durable record of
what was actually paid against a bill.
"""
from __future__ import annotations

import csv
import datetime
import io
import zipfile

from . import data_schema as S
from .canonicalizer import _canonical_bill, _canonical_payment, _canonical_tenant, _recompute
from .csv_parser import sheet_rows_to_csv


# ---------------------------------------------------------------------------
# Canonical model from the database
# ---------------------------------------------------------------------------


def _selected_ids(tenants_filter) -> set[int] | None:
    """Resolve a `tenants_list` query value into a set of tenant ids (None = all)."""
    if tenants_filter in (None, "", "*", "all"):
        return None
    ids: set[int] = set()
    for part in str(tenants_filter).split(","):
        part = part.strip()
        if part.isdigit():
            ids.add(int(part))
    return ids or None


def _decrypted_pins(tenant_ids, landlord_id=None) -> dict[int, str]:
    """Map tenantId -> plain PIN, scoped to the tenants being exported."""
    from app.authentication.common.pin_vault import decrypt_admin_view_pin
    from app.core.db import get_conn

    if not tenant_ids:
        return {}
    pins: dict[int, str] = {}
    try:
        with get_conn() as conn:
            rows = conn.execute(
                'SELECT "tenantId", \"encryptedPin\" FROM \"tenantPinAdminStore\" ORDER BY "tenantId"'
            ).fetchall()
    except Exception:
        return pins
    for row in rows:
        try:
            tid = int(row["tenantId"] or 0)
        except (TypeError, ValueError):
            continue
        if tid not in tenant_ids:
            continue
        try:
            pins[tid] = decrypt_admin_view_pin(row["encryptedPin"]) or ""
        except Exception:
            pins[tid] = ""
    return pins


def _fetch_payments(tenant_ids, landlord_id=None) -> list[dict]:
    """All ACTIVE payment_entries for the selected tenants (canonical inputs)."""
    from app.core.db import get_conn

    if not tenant_ids:
        return []
    try:
        with get_conn() as conn:
            rows = conn.execute(
                "SELECT * FROM \"paymentEntries\" "
                '''WHERE status = 'ACTIVE' AND "tenantId" = ANY(%s) '''
                'ORDER BY "tenantId", "billNo", id',
                (sorted(tenant_ids),),
            ).fetchall()
    except Exception:
        return []
    return [dict(r) for r in rows]


def build_canonical(landlord_id=None, tenants_filter=None) -> dict:
    """Export the canonical model: {tenants, bills, payments, warnings, ...}."""
    from app.services.billing_service import get_all_receipts
    from app.services.tenant_service import load_tenants

    tenants = load_tenants(include_archived=True, landlord_id=landlord_id)
    receipts = get_all_receipts(include_archived_tenants=True, landlord_id=landlord_id)

    selected_ids = _selected_ids(tenants_filter)
    if selected_ids is not None:
        tenants = [t for t in tenants if t.id in selected_ids]
        receipts = [
            r for r in receipts
            if r.get("TenantId") is not None and int(r.get("TenantId") or 0) in selected_ids
        ]
    tenant_ids = {t.id for t in tenants}
    pins = _decrypted_pins(tenant_ids, landlord_id)

    tenants_out: dict[int, dict] = {}
    for t in tenants:
        num_id, profile = _canonical_tenant(
            {
                "tenantId": S.format_tenant_id(t.id),
                "tenantName": getattr(t, "name", "") or "",
                "Phone": getattr(t, "phone", "") or "",
                "Email": getattr(t, "email", "") or "",
                "Company": getattr(t, "company", "") or "",
                "Address": getattr(t, "address", "") or "",
                "Room": getattr(t, "roomNumber", "") or "",
                "meterId": getattr(t, "meterId", "") or "",
                "PIN": pins.get(t.id, ""),
                "Rent": getattr(t, "rent", 0) or 0,
                "Water": getattr(t, "water", 0) or 0,
                "electricityRate": getattr(t, "electricityRate", 0) or 0,
                "additionalPersonRate": getattr(t, "additionalPersonCharge", 0) or 0,
                "tankWater": getattr(t, "defaulttankWaterCharge", 0) or 0,
                "Status": getattr(t, "status", "Active") or "Active",
            }
        )
        if num_id > 0:
            tenants_out.setdefault(num_id, profile)

    bills_out: dict[tuple[int, str], dict] = {}
    stored_amount: dict[tuple[int, str], float] = {}
    for r in receipts:
        num_id = r.get("TenantId")
        try:
            num_id = int(num_id or 0)
        except (TypeError, ValueError):
            continue
        if num_id not in tenant_ids:
            continue
        bill_no = S.clean(r.get("Bill"))
        if not bill_no:
            continue
        bills_out.setdefault(
            (num_id, bill_no),
            _canonical_bill(
                num_id,
                {
                    "BillNo": bill_no,
                    "Month": r.get("Month", ""),
                    "Date": r.get("Date", ""),
                    "Previous": r.get("Previous", 0),
                    "Current": r.get("Current", 0),
                    "Units": r.get("Units", 0),
                    "Rent": r.get("Rent", 0),
                    "Water": r.get("Water", 0),
                    "Electricity": r.get("Electricity", 0),
                    "Additional": r.get("Additional", 0),
                    "tankWater": r.get("tankWater", 0),
                    "Maintenance": r.get("MaintenanceCharge", 0),
                    "MaintenanceDesc": r.get("MaintenanceDesc", ""),
                    "Arrears": r.get("previousArrears", 0),
                    "amountReceived": r.get("amountReceived", 0),
                    "paymentStatus": r.get("paymentStatus", ""),
                    "receiptStatus": r.get("Status", "ACTIVE") or "ACTIVE",
                    "Total": r.get("Total", 0),
                },
            ),
        )
        stored_amount.setdefault((num_id, bill_no), S.to_float(r.get("amountReceived")))

    payments_out: list[dict] = []
    seen_external: set[str] = set()
    seq_counter: dict[tuple[int, str], int] = {}

    for row in _fetch_payments(tenant_ids, landlord_id):
        tid = row.get("tenantId")
        try:
            tid = int(tid or 0)
        except (TypeError, ValueError):
            continue
        bill_no = S.clean(row.get("billNo"))
        bill = bills_out.get((tid, bill_no))
        if tid not in tenants_out or bill is None or not bill_no:
            continue
        external_id = S.clean(row.get("externalId"))
        seq_counter[(tid, bill_no)] = seq_counter.get((tid, bill_no), 0) + 1
        payment = _canonical_payment(
            {
                "tenantId": S.format_tenant_id(tid),
                "billNo": bill_no,
                "paymentId": external_id,
                "paymentDate": row.get("paymentDate"),
                "paymentAmount": row.get("paymentAmount"),
                "paymentMethod": row.get("paymentMethod") or "OTHER",
                "paymentReference": row.get("reference"),
                "paymentNotes": row.get("notes"),
                "paymentSource": row.get("source"),
            },
            bill,
            seq_counter[(tid, bill_no)],
            explicit_id=bool(external_id),
        )
        if payment["amount"] <= 0 or payment["externalId"] in seen_external:
            continue
        seen_external.add(payment["externalId"])
        payments_out.append(payment)

    # Legacy backfill: bills with a stored v1 amount but no durable payment row
    # synthesise a single LEGACY_IMPORT payment so derived fields survive export.
    legacy_count = 0
    paid_keys = {(p["tenantId"], p["billNo"]) for p in payments_out}
    for (num_id, bill_no), bill in bills_out.items():
        amount = stored_amount.get((num_id, bill_no), 0)
        if amount <= 0 or (num_id, bill_no) in paid_keys:
            continue
        seq_counter[(num_id, bill_no)] = seq_counter.get((num_id, bill_no), 0) + 1
        payment = _canonical_payment(
            {
                "tenantId": S.format_tenant_id(num_id),
                "billNo": bill_no,
                "paymentId": "",
                "paymentDate": bill.get("Date"),
                "paymentAmount": amount,
                "paymentMethod": "OTHER",
                "paymentReference": "",
                "paymentNotes": "Legacy v1 receipt amount",
                "paymentSource": "LEGACY_IMPORT",
            },
            bill,
            seq_counter[(num_id, bill_no)],
            explicit_id=False,
        )
        if payment["amount"] <= 0 or payment["externalId"] in seen_external:
            continue
        seen_external.add(payment["externalId"])
        payments_out.append(payment)
        legacy_count += 1

    result = {
        "tenants": tenants_out,
        "bills": list(bills_out.values()),
        "payments": payments_out,
        "warnings": [],
        "version": "v2",
        "legacy_backfilled": legacy_count,
    }
    _recompute(result)
    return result


# ---------------------------------------------------------------------------
# Flat 40-column CSV
# ---------------------------------------------------------------------------


def _flat_row(tenants, bill: dict, payment: dict | None) -> dict:
    row = dict(S.BLANK_FLAT_ROW)
    profile = tenants.get(bill["tenantId"]) or {}
    for h in S.PROFILE_HEADERS:
        v = profile.get(h, "")
        if v not in (None, ""):
            row[h] = v
    row["tenantId"] = S.format_tenant_id(bill["tenantId"])
    for h in S.BILL_FLAT_HEADERS:
        v = bill.get(h, "")
        if v not in (None, ""):
            row[h] = v
    if payment is not None:
        row["paymentId"] = payment.get("externalId") or ""
        row["paymentDate"] = payment.get("paymentDate") or ""
        row["paymentAmount"] = payment.get("amount", 0)
        row["paymentMethod"] = payment.get("paymentMethod") or ""
        row["paymentReference"] = payment.get("reference") or ""
        row["paymentNotes"] = payment.get("notes") or ""
        row["paymentSource"] = payment.get("source") or ""
    return row


def flat_rows(canonical: dict) -> list[dict]:
    """One row per payment, grouped by bill; bills without payments get one blank row."""
    tenants = canonical["tenants"]
    grouped: dict[tuple[int, str], list[dict]] = {}
    for p in canonical["payments"]:
        grouped.setdefault((p["tenantId"], p["billNo"]), []).append(p)
    rows: list[dict] = []
    for bill in canonical["bills"]:
        pays = grouped.get((bill["tenantId"], bill["billNo"]))
        if pays:
            rows.extend(_flat_row(tenants, bill, p) for p in pays)
        else:
            rows.append(_flat_row(tenants, bill, None))
    return rows


def flat_csv_bytes(canonical: dict) -> bytes:
    return sheet_rows_to_csv(list(S.FLAT_HEADERS), flat_rows(canonical))


# ---------------------------------------------------------------------------
# 3-sheet XLSX
# ---------------------------------------------------------------------------


def _receipt_row(bill: dict) -> dict:
    return {
        "BillNo": bill.get("billNo", ""),
        "tenantId": S.format_tenant_id(bill.get("tenantId", 0) or 0),
        "Month": bill.get("Month", ""),
        "Date": bill.get("Date", ""),
        "Previous": bill.get("Previous", 0),
        "Current": bill.get("Current", 0),
        "Units": bill.get("Units", 0),
        "Rent": bill.get("Rent", 0),
        "Water": bill.get("Water", 0),
        "Electricity": bill.get("Electricity", 0),
        "Additional": bill.get("Additional", 0),
        "tankWater": bill.get("tankWater", 0),
        "Maintenance": bill.get("Maintenance", 0),
        "Arrears": bill.get("Arrears", 0),
        "amountReceived": bill.get("amountReceived", 0),
        "Total": bill.get("Total", 0),
        "paymentStatus": bill.get("paymentStatus", "PENDING"),
        "receiptStatus": bill.get("receiptStatus", "ACTIVE"),
    }


def _payment_row(payment: dict) -> dict:
    return {
        "paymentId": payment.get("externalId") or "",
        "billNo": payment.get("billNo", ""),
        "tenantId": S.format_tenant_id(payment.get("tenantId", 0) or 0),
        "paymentDate": payment.get("paymentDate", ""),
        "paymentAmount": payment.get("amount", 0),
        "paymentMethod": payment.get("paymentMethod", ""),
        "paymentReference": payment.get("reference") or "",
        "paymentNotes": payment.get("notes") or "",
        "paymentStatus": "",
        "paymentSource": payment.get("source", "IMPORT"),
        "paymentCreatedBy": "",
    }


def workbook(canonical: dict):
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    tenants = canonical["tenants"]
    wb = openpyxl.Workbook()
    ws_profile = wb.active
    ws_profile.title = S.SHEET_PROFILE
    ws_receipts = wb.create_sheet(S.SHEET_RECEIPTS)
    ws_payments = wb.create_sheet(S.SHEET_PAYMENTS)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F81BD")

    profile_rows = [tenants[i] for i in sorted(tenants)]
    receipt_rows = [_receipt_row(b) for b in canonical["bills"]]
    payment_rows = [_payment_row(p) for p in canonical["payments"]]

    for ws, headers, rows in (
        (ws_profile, S.PROFILE_HEADERS, profile_rows),
        (ws_receipts, S.RECEIPT_HEADERS, receipt_rows),
        (ws_payments, S.PAYMENT_HEADERS, payment_rows),
    ):
        ws.append(list(headers))
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        for row in rows:
            ws.append([row.get(h, "") for h in headers])
    return wb


def workbook_bytes(canonical: dict) -> bytes:
    buf = io.BytesIO()
    workbook(canonical).save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Full ZIP (v1 parity + canonical artifacts)
# ---------------------------------------------------------------------------


def full_zip_bytes(landlord_id=None, tenants_filter=None) -> bytes:
    from app.services.billing_service import get_all_receipts
    from app.services.landlord_config_service import get_effective_landlord_config
    from app.services.pdf_service import generate_professional_pdf
    from app.services.tenant_service import load_tenants

    canonical = build_canonical(landlord_id, tenants_filter)
    tenants = load_tenants(include_archived=True, landlord_id=landlord_id)
    receipts = get_all_receipts(include_archived_tenants=True, landlord_id=landlord_id)

    selected_ids = _selected_ids(tenants_filter)
    if selected_ids is not None:
        tenants = [t for t in tenants if t.id in selected_ids]
        receipts = [
            r for r in receipts
            if r.get("TenantId") is not None and int(r.get("TenantId") or 0) in selected_ids
        ]

    landlord_conf = get_effective_landlord_config(landlord_id) if landlord_id else None

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        stream = io.StringIO()
        if receipts:
            writer = csv.DictWriter(stream, fieldnames=list(receipts[0].keys()))
            writer.writeheader()
            writer.writerows(receipts)
        zf.writestr("receipts_data.csv", stream.getvalue())

        zf.writestr("flat_import_40col.csv", flat_csv_bytes(canonical))
        zf.writestr("Rent_Data_Export_V2.xlsx", workbook_bytes(canonical))

        for r in receipts:
            try:
                formatted_date = datetime.datetime.strptime(
                    r.get("Date", ""), "%d %B %Y"
                ).strftime("%Y%m%d")
            except Exception:
                formatted_date = str(r.get("Date", "")).replace(" ", "")
            tenant_name = str(r.get("Tenant", "Unknown")).replace(" ", "_")
            filename = f"{tenant_name}_{formatted_date}_{r['Bill']}.pdf"
            folder = "archive" if str(r.get("Status", "ACTIVE")).upper() == "ARCHIVED" else "active"
            try:
                pdf_stream = generate_professional_pdf(r, landlord_conf)
                zf.writestr(f"PDFs/{folder}/{filename}", pdf_stream.getvalue())
            except Exception as exc:
                zf.writestr(f"PDFs/FAILED_{filename}.txt", f"PDF generation failed: {exc}")
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Endpoint helper
# ---------------------------------------------------------------------------


def export_bytes(export_format="xlsx", landlord_id=None, tenants_filter=None):
    """Return (content_bytes, filename, media_type) for the v2 export endpoint.

    Raises ValueError for anything other than csv / xlsx / zip.
    """
    fmt = (export_format or "xlsx").lower()
    date_suffix = datetime.date.today().strftime("%Y%m%d")
    if fmt == "csv":
        data = flat_csv_bytes(build_canonical(landlord_id, tenants_filter))
        return data, f"flat_import_40col_{date_suffix}.csv", "text/csv"
    if fmt == "xlsx":
        data = workbook_bytes(build_canonical(landlord_id, tenants_filter))
        return (
            data,
            f"Rent_Data_Export_V2_{date_suffix}.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    if fmt == "zip":
        return (
            full_zip_bytes(landlord_id, tenants_filter),
            f"tenants_data_v2_{date_suffix}.zip",
            "application/zip",
        )
    raise ValueError("Unsupported format. Use 'xlsx', 'zip', or 'csv'.")


# ---------------------------------------------------------------------------
# Templates (Phase H)
# ---------------------------------------------------------------------------


def template_canonical() -> dict:
    """Small self-consistent sample that demonstrates the full 40-col layout."""
    sample = {
        "tenants": {
            1: {
                "tenantId": "T001",
                "tenantName": "John Doe",
                "Phone": "9876543210",
                "Email": "john@gmail.com",
                "Company": "Acme Pvt Ltd",
                "Address": "Delhi",
                "Room": "A101",
                "meterId": "MTR001",
                "PIN": "",
                "Rent": 15000.0,
                "Water": 500.0,
                "electricityRate": 8.5,
                "additionalPersonRate": 1000.0,
                "tankWater": 300.0,
                "Status": "Active",
            },
            2: {
                "tenantId": "T002",
                "tenantName": "Jane Smith",
                "Phone": "9123456789",
                "Email": "jane@gmail.com",
                "Company": "XYZ Solutions",
                "Address": "Mumbai",
                "Room": "B202",
                "meterId": "MTR002",
                "PIN": "",
                "Rent": 20000.0,
                "Water": 750.0,
                "electricityRate": 8.5,
                "additionalPersonRate": 1500.0,
                "tankWater": 400.0,
                "Status": "Active",
            },
        },
        "bills": [
            {
                "tenantId": 1,
                "billNo": "T1-001",
                "Month": "January 2026",
                "Date": "01 January 2026",
                "Previous": 120.0,
                "Current": 150.0,
                "Units": 30.0,
                "Rent": 15000.0,
                "Water": 500.0,
                "Electricity": 255.0,
                "Additional": 1000.0,
                "tankWater": 300.0,
                "Maintenance": 0.0,
                "MaintenanceDesc": "",
                "Arrears": 0.0,
                "amountReceived": 17055.0,
                "paymentStatus": "PAID",
                "receiptStatus": "ACTIVE",
                "Total": 17055.0,
                "_has_total": True,
            },
            {
                "tenantId": 2,
                "billNo": "T2-001",
                "Month": "January 2026",
                "Date": "01 January 2026",
                "Previous": 0.0,
                "Current": 40.0,
                "Units": 40.0,
                "Rent": 20000.0,
                "Water": 750.0,
                "Electricity": 340.0,
                "Additional": 1500.0,
                "tankWater": 400.0,
                "Maintenance": 0.0,
                "MaintenanceDesc": "",
                "Arrears": 0.0,
                "amountReceived": 0.0,
                "paymentStatus": "PENDING",
                "receiptStatus": "ACTIVE",
                "Total": 22990.0,
                "_has_total": True,
            },
        ],
        "payments": [
            {
                "tenantId": 1,
                "billNo": "T1-001",
                "paymentDate": "2026-01-31",
                "amount": 17055.0,
                "paymentMethod": "CASH",
                "reference": "RC-001",
                "notes": "January rent paid",
                "externalId": "PAY-1-T1-001-00",
                "source": "IMPORT",
            }
        ],
        "warnings": [],
        "version": "v2",
        "legacy_backfilled": 0,
    }
    _recompute(sample)
    return sample


def template_csv_bytes() -> bytes:
    return flat_csv_bytes(template_canonical())


def template_workbook_bytes() -> bytes:
    return workbook_bytes(template_canonical())